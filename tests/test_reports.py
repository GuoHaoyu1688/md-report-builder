import importlib.util
import json
import zipfile
from pathlib import Path

import pytest
from bs4 import BeautifulSoup
from docx import Document
from docx.oxml.ns import qn
from lxml import etree
from PIL import Image
from pypdf import PdfReader

from md_report_builder import build
from md_report_builder.core.errors import ReportError
from md_report_builder.core.history import History
from md_report_builder.core.model import chart_image, chinese_number, heading_number, math_image
from md_report_builder.core.source import lint_markdown, read_source, yaml_load
from md_report_builder.core.styles import load_template, points

ROOT = Path(__file__).resolve().parents[1]
EXAMPLE = ROOT / "examples/report.md"
NS = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}


@pytest.mark.parametrize("template", ["business", "gongwen", "tech", "weekly"])
def test_all_templates_docx_and_html_with_source_assertions(tmp_path, template):
    result = build(EXAMPLE, "docx", template, tmp_path)
    assert result["counts"] == {"tables": 3, "figures": 2, "headings": 6}
    document = Document(result["output"])
    assert len(document.tables) == 3
    assert document.tables[0].cell(4, 2).text == "96"
    assert document.tables[1].cell(3, 2).text == "96"
    assert document.tables[2].cell(3, 2).text == "Complete"
    assert document.sections[0].header.paragraphs[0].text == "Example operations review"
    text = "\n".join(p.text for p in document.paragraphs)
    assert "Figure 1" in text and "Figure 2" in text and "Table 3" in text
    assert "2026-01-15" in text and "Example Analyst" in text
    assert any(p.style.name == "Heading 3" for p in document.paragraphs)
    assert not document.styles["Title"].element.xpath("./w:pPr/w:pBdr")
    assert document.styles["Normal"].element.rPr.rFonts.get(qn("w:eastAsia")) == "Noto Serif CJK SC"
    with zipfile.ZipFile(result["output"]) as archive:
        xml = etree.fromstring(archive.read("word/document.xml"))
        assert len(xml.xpath("//w:tblHeader", namespaces=NS)) == 3
        assert len(xml.xpath("//w:footnoteReference", namespaces=NS)) == 1
        instructions = " ".join(xml.xpath("//w:instrText/text()", namespaces=NS))
        if template != "weekly":
            assert "TOC" in instructions
        assert "SEQ Figure" in instructions and "REF fig_volume" in instructions
        assert len(xml.xpath("//w:bookmarkStart", namespaces=NS)) == 11
        footnotes = etree.fromstring(archive.read("word/footnotes.xml"))
        assert "Counts are fictional." in "".join(footnotes.itertext())
        footer = archive.read("word/footer1.xml").decode()
        assert "NUMPAGES" in footer and "PAGE" in footer
    html = build(EXAMPLE, "html", template, tmp_path)
    text = Path(html["output"]).read_text()
    soup = BeautifulSoup(text, "html.parser")
    assert len(soup.select("table")) == 3
    assert len(soup.select("figure")) == 2
    assert all(img["src"].startswith("data:image/png;base64,") for img in soup.select("img"))
    assert soup.select_one("#fig-volume")
    assert soup.select_one('a[href="#fig-volume"]').text == "Figure 1"
    assert "Counts are fictional" in soup.select_one("#fn-1").text
    assert not soup.select("script,link[rel=stylesheet]")
    assert "[@" not in soup.get_text()


@pytest.mark.skipif(
    importlib.util.find_spec("weasyprint") is None, reason="PDF extra not installed"
)
@pytest.mark.parametrize("template", ["business", "gongwen", "tech", "weekly"])
def test_real_pdf_all_templates(tmp_path, template):
    result = build(EXAMPLE, "pdf", template, tmp_path)
    document = PdfReader(result["output"])
    text = "\n".join(page.extract_text() for page in document.pages)
    assert 2 <= len(document.pages) <= 10
    for required in [
        "Example operations review",
        "96",
        "Figure 1",
        "Figure 2",
        "Table 3",
        "Counts are fictional",
        "Next actions",
    ]:
        assert required in text
    assert document.outline
    assert "Page 1 of" in document.pages[0].extract_text()
    assert any(page.images for page in document.pages)
    assert any(page.get("/Annots") for page in document.pages)
    for page in document.pages:
        assert len(page.extract_text().strip()) > 15


def test_numbering_and_font_sizes_hand_checked():
    counters = [0] * 6
    assert heading_number(counters, 1, "technical") == "1 "
    assert heading_number(counters, 2, "technical") == "1.1 "
    assert heading_number(counters, 3, "technical") == "1.1.1 "
    assert heading_number(counters, 2, "technical") == "1.2 "
    assert heading_number(counters, 1, "technical") == "2 "
    counters = [0] * 6
    assert [heading_number(counters, i, "gongwen") for i in range(1, 5)] == [
        "一、",
        "（一）",
        "1.",
        "（1）",
    ]
    assert [chinese_number(n) for n in [10, 11, 20, 99, 100]] == [
        "十",
        "十一",
        "二十",
        "九十九",
        "100",
    ]
    assert points("三号") == 16 and points("四号") == 14 and points("小四") == 12
    assert points(10.5) == 10.5
    with pytest.raises(ReportError):
        points("unknown")


def test_directory_order_includes_variables_and_unicode(tmp_path):
    source = tmp_path / "source"
    source.mkdir()
    (source / "02.md").write_text("# Second\n\nSecond section", encoding="utf-8")
    (source / "01.md").write_text(
        "---\ntitle: 示例报告\nauthor: ''\ndate: 2026-01-15\nlanguage: zh\n---\n"
        "# 示例\n\n{{date}} 项目数量\n\n!include parts/tail.md\n",
        encoding="utf-8",
    )
    (source / "parts").mkdir()
    (source / "parts/tail.md").write_text("## 结论\n\n示例结果：完成两项。", encoding="utf-8")
    read = read_source(source)
    assert read["markdown"].index("示例") < read["markdown"].index("Second")
    assert "2026-01-15 项目数量" in read["markdown"]
    result = build(source, "docx", "gongwen", tmp_path / "out")
    document = Document(result["output"])
    text = "\n".join(p.text for p in document.paragraphs)
    assert "一、示例" in text and "（一）结论" in text
    assert "第 " in document.sections[0].footer.paragraphs[0].text
    if importlib.util.find_spec("weasyprint"):
        pdf = build(source, "pdf", "gongwen", tmp_path / "out")
        text = "".join(page.extract_text() for page in PdfReader(pdf["output"]).pages)
        assert "示例报告" in text and "完成两项" in text


def test_lint_four_required_categories(tmp_path):
    source = tmp_path / "bad.md"
    source.write_text(
        "# Start\n\n### Skipped\n\n[empty]()\n\n![missing](missing.png)\n\n"
        "| A | B |\n| --- | --- |\n| 1 | 2 | 3 |\n"
    )
    result = lint_markdown(source)
    assert not result["valid"]
    assert {d["code"] for d in result["diagnostics"]} == {
        "heading_jump",
        "empty_link",
        "missing_image",
        "table_columns",
    }
    with pytest.raises(ReportError, match="lint diagnostics"):
        build(source, "html", out=tmp_path / "out", strict=True)


@pytest.mark.parametrize(
    "content,code",
    [
        ("!include a.md", "include_cycle"),
        ("!include ../outside.md", "path_escape"),
        ("!include https://example.com/a.md", "network_disabled"),
        ("---\ntitle: A\ntitle: B\n---\n# Body", "invalid_yaml"),
        ("---\ntitle: A\n# Body", "invalid_front_matter"),
        ("# Heading\n\nSee [@fig:missing]", "unknown_reference"),
        ("# First {#same}\n# Second {#same}", "duplicate_reference"),
    ],
)
def test_rejects_unsafe_or_ambiguous_source(tmp_path, content, code):
    path = tmp_path / "a.md"
    path.write_text(content)
    with pytest.raises(ReportError) as error:
        build(path, "html", out=tmp_path / "out")
    assert error.value.code == code


def test_images_are_embedded_and_remote_images_blocked(tmp_path, monkeypatch):
    Image.new("RGB", (100, 50), "teal").save(tmp_path / "local.png")
    source = tmp_path / "image.md"
    source.write_text(
        "# Images\n\n![Local chart {#local}](local.png)\n\n**Strong** *emphasis* `code` and [link](https://example.com).\n"
    )
    import socket

    monkeypatch.setattr(
        socket, "create_connection", lambda *a, **k: pytest.fail("Network attempted")
    )
    result = build(source, "html", out=tmp_path / "out")
    text = Path(result["output"]).read_text()
    assert "data:image/png;base64," in text and "<strong>Strong</strong>" in text
    assert 'id="fig-local"' in text
    result = build(source, "docx", out=tmp_path / "out")
    assert len(Document(result["output"]).inline_shapes) == 1
    source.write_text("![Remote](https://example.com/image.png)")
    with pytest.raises(ReportError) as error:
        build(source, "html", out=tmp_path / "out")
    assert error.value.code == "network_disabled"


def test_code_is_not_expanded_and_mermaid_is_preserved(tmp_path):
    source = tmp_path / "code.md"
    source.write_text(
        "```sh\n!include missing.md\n```\n\n```mermaid\ngraph LR\n A-->B\n```\n\n{{unknown}}"
    )
    result = build(source, "html", out=tmp_path / "out")
    assert len(result["warnings"]) == 2
    text = Path(result["output"]).read_text()
    assert "!include missing.md" in text and "graph LR" in text and "{{unknown}}" in text


def test_chart_csv_yaml_math_and_invalid_values():
    data, spec = chart_image("Group,Count\nA,2\nB,3\n")
    assert data.startswith(b"\x89PNG") and spec["x"] == ["A", "B"]
    data, _ = chart_image("type: pie\nx: [A, B]\ny: [2, 3]")
    assert data.startswith(b"\x89PNG")
    assert math_image(r"\frac{1}{2} + x^2").startswith(b"\x89PNG")
    with pytest.raises(ReportError):
        chart_image("type: bar\nx: [A, B]\ny: [1]")
    with pytest.raises(ReportError):
        chart_image("type: bar\nx: [A]\ny: [.nan]")
    with pytest.raises(ReportError):
        math_image(r"\notarealcommand{x}")
    with pytest.raises(ReportError):
        yaml_load("x: !!python/object:foo {}")


def test_custom_template_reference_and_atomic_failure(tmp_path):
    template = tmp_path / "template"
    template.mkdir()
    (template / "style.yaml").write_text(
        "name: custom\nbody_size: 小四\nnumbering: none\n", encoding="utf-8"
    )
    reference = Document()
    reference.add_paragraph("Stale reference body")
    reference.save(template / "reference.docx")
    source = tmp_path / "source.md"
    source.write_text("# A\n\nHello")
    result = build(source, "docx", str(template), tmp_path / "out")
    assert "Stale reference body" not in "\n".join(
        p.text for p in Document(result["output"]).paragraphs
    )
    before = Path(result["output"]).read_bytes()
    source.write_text("# A\n\n![Missing](absent.png)")
    with pytest.raises(ReportError):
        build(source, "docx", str(template), tmp_path / "out")
    assert Path(result["output"]).read_bytes() == before
    assert not list((tmp_path / "out").glob(".mdr-*"))
    (template / "style.yaml").write_text("name: custom\nmargin_left: 90")
    with pytest.raises(ReportError):
        load_template(str(template))


@pytest.mark.parametrize("format", ["json", "csv"])
def test_history_roundtrip_idempotency_and_backup(tmp_path, format):
    source = tmp_path / "source.md"
    source.write_text("# Example\n\nHello")
    result = build(source, "html", "weekly", tmp_path / "out", data_dir=tmp_path / "data")
    second = build(source, "html", "weekly", tmp_path / "out", data_dir=tmp_path / "data")
    assert result["id"] == second["id"] and result["sha256"] == second["sha256"]
    store = History(tmp_path / "data")
    assert len(store.list()) == 1
    snapshot = tmp_path / f"history.{format}"
    store.export(snapshot, format)
    if format == "csv":
        assert snapshot.read_bytes().startswith(b"\xef\xbb\xbf")
    restored = History(tmp_path / "restored")
    assert restored.import_file(snapshot) == {"imported": 1, "skipped": 0}
    assert restored.import_file(snapshot) == {"imported": 0, "skipped": 1}
    assert restored.list() == store.list()
    with zipfile.ZipFile(store.backup()["path"]) as archive:
        assert archive.namelist() == ["cache.sqlite"]
    with pytest.raises(ReportError):
        store.clear()
    assert store.clear(True) == {"deleted": 1}
    invalid = tmp_path / "invalid.json"
    invalid.write_text(json.dumps({"builds": [second, {"id": "bad"}]}))
    with pytest.raises(ReportError):
        store.import_file(invalid)
    assert store.list() == []


def test_empty_and_uneven_chart_csv_fail_cleanly():
    for content in ["", "Group,Count\nA\n", "type: bar\nx: [A]\nseries: {}"]:
        with pytest.raises(ReportError):
            chart_image(content)
    data, _ = chart_image("# A YAML comment\ntype: bar\nx: [A]\ny: [2]")
    assert data.startswith(b"\x89PNG")


def test_chinese_chart_labels_use_a_local_font():
    data, _ = chart_image("type: bar\nx: [示例甲, 示例乙]\ny: [2, 3]")
    assert data.startswith(b"\x89PNG")


def test_code_table_is_not_linted_and_directory_symlink_is_rejected(tmp_path):
    source = tmp_path / "source"
    source.mkdir()
    (source / "example.md").write_text("```md\n| A | B |\n| --- | --- |\n| 1 | 2 | 3 |\n```\n")
    assert lint_markdown(source)["valid"]
    outside = tmp_path / "outside.md"
    outside.write_text("# Outside")
    (source / "link.md").symlink_to(outside)
    with pytest.raises(ReportError) as error:
        read_source(source)
    assert error.value.code == "path_escape"
