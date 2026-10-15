"""Native python-docx report renderer with fields, bookmarks and true footnotes."""

from __future__ import annotations

import io
import re

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.opc.packuri import PackURI
from docx.opc.part import Part
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor
from lxml import etree
from PIL import Image

from .model import image_bytes, math_image


def field(paragraph, instruction: str, cached: str):
    for kind in ("begin", "instruction", "separate", "value", "end"):
        run = paragraph.add_run()
        if kind in {"begin", "separate", "end"}:
            element = OxmlElement("w:fldChar")
            element.set(qn("w:fldCharType"), kind)
        elif kind == "instruction":
            element = OxmlElement("w:instrText")
            element.set(qn("xml:space"), "preserve")
            element.text = " " + instruction + " "
        else:
            run.text = cached
            continue
        run._r.append(element)


def bookmark(paragraph, name, identifier):
    start = OxmlElement("w:bookmarkStart")
    start.set(qn("w:id"), str(identifier))
    start.set(qn("w:name"), name.replace("-", "_"))
    end = OxmlElement("w:bookmarkEnd")
    end.set(qn("w:id"), str(identifier))
    paragraph._p.insert(0, start)
    paragraph._p.append(end)


def hyperlink(paragraph, target: str, label: str, internal=False):
    element = OxmlElement("w:hyperlink")
    if internal:
        element.set(qn("w:anchor"), target.lstrip("#").replace("-", "_"))
    else:
        relation = paragraph.part.relate_to(target, RT.HYPERLINK, is_external=True)
        element.set(qn("r:id"), relation)
    run = OxmlElement("w:r")
    props = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "215B70")
    props.append(color)
    run.append(props)
    text = OxmlElement("w:t")
    text.text = label
    run.append(text)
    element.append(run)
    paragraph._p.append(element)


def add_image(paragraph, data: bytes, width_mm: float, inline=False, alt=""):
    with Image.open(io.BytesIO(data)) as image:
        width, height = image.size
    natural_mm = width / (220 if inline else 180) * 25.4
    display_width = min(width_mm, natural_mm)
    if inline:
        display_width = 4.2 * width / height
    shape = paragraph.add_run().add_picture(io.BytesIO(data), width=Mm(display_width))
    shape._inline.docPr.set("descr", alt)
    return shape


def inline_docx(paragraph, children, model, width_mm):
    bold = italic = strike = False
    current_link = None
    for token in children:
        kind = token.type
        if kind == "strong_open":
            bold = True
        elif kind == "strong_close":
            bold = False
        elif kind == "em_open":
            italic = True
        elif kind == "em_close":
            italic = False
        elif kind == "s_open":
            strike = True
        elif kind == "s_close":
            strike = False
        elif kind == "link_open":
            current_link = token.attrGet("href")
        elif kind == "link_close":
            current_link = None
        elif kind in {"softbreak", "hardbreak"}:
            paragraph.add_run("\n" if kind == "hardbreak" else " ")
        elif kind == "footnote_ref":
            reference = OxmlElement("w:footnoteReference")
            reference.set(qn("w:id"), str(token.meta["id"] + 1))
            run = paragraph.add_run()
            run.font.superscript = True
            run._r.append(reference)
        elif kind == "math_inline":
            add_image(
                paragraph, math_image(token.content), width_mm, inline=True, alt=token.content
            )
        elif kind == "image":
            add_image(
                paragraph,
                image_bytes(token.attrGet("src"), model["source"]),
                width_mm,
                alt=token.content,
            )
        else:
            content = re.sub(
                r"^\[([ xX])\]\s", lambda m: "☑ " if m[1].lower() == "x" else "☐ ", token.content
            )
            if current_link:
                hyperlink(paragraph, current_link, content, internal=current_link.startswith("#"))
                continue
            for piece in re.split(r"(\[@(?:sec|fig|tbl):[a-zA-Z][a-zA-Z0-9_-]*\])", content):
                if piece.startswith("[@"):
                    ref = model["refs"][piece[2:-1]]
                    field(paragraph, f"REF {ref['anchor'].replace('-', '_')} \\h", ref["label"])
                elif piece:
                    run = paragraph.add_run(piece)
                    run.bold, run.italic, run.font.strike = bold, italic, strike
                    if kind == "code_inline":
                        run.font.name = "Courier New"
                        run.font.size = Pt(10)


def add_footnotes(document, model):
    if not model["footnotes"]:
        return
    root = OxmlElement("w:footnotes")
    for identifier, kind in [(-1, "separator"), (0, "continuationSeparator")]:
        note = OxmlElement("w:footnote")
        note.set(qn("w:type"), kind)
        note.set(qn("w:id"), str(identifier))
        paragraph, run = OxmlElement("w:p"), OxmlElement("w:r")
        run.append(OxmlElement(f"w:{kind}"))
        paragraph.append(run)
        note.append(paragraph)
        root.append(note)
    for identifier, children in model["footnotes"].items():
        note = OxmlElement("w:footnote")
        note.set(qn("w:id"), str(identifier + 1))
        paragraph = OxmlElement("w:p")
        props = OxmlElement("w:pPr")
        style = OxmlElement("w:pStyle")
        style.set(qn("w:val"), "FootnoteText")
        props.append(style)
        paragraph.append(props)
        run = OxmlElement("w:r")
        run.append(OxmlElement("w:footnoteRef"))
        paragraph.append(run)
        run = OxmlElement("w:r")
        text = OxmlElement("w:t")
        text.set(qn("xml:space"), "preserve")
        text.text = " " + "".join(t.content for t in children)
        run.append(text)
        paragraph.append(run)
        note.append(paragraph)
        root.append(note)
    part = Part(
        PackURI("/word/footnotes.xml"),
        "application/vnd.openxmlformats-officedocument.wordprocessingml.footnotes+xml",
        etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True),
        document.part.package,
    )
    document.part.relate_to(part, RT.FOOTNOTES)


def write_docx(model: dict, path):
    style, meta = model["style"], model["metadata"]
    document = Document(style["reference"]) if style["reference"] else Document()
    # A reference contributes styles and sections, not stale body content.
    for child in list(document.element.body):
        if child.tag != qn("w:sectPr"):
            document.element.body.remove(child)
    document.core_properties.author = meta.get("author", "")
    document.core_properties.title = meta["title"]
    document.core_properties.subject = meta.get("subtitle", "")
    document.core_properties.last_modified_by = "md-report-builder"
    for name in (
        "Normal",
        "Title",
        "Subtitle",
        "Caption",
        "Header",
        "Footer",
        *[f"Heading {i}" for i in range(1, 7)],
    ):
        item = document.styles[name]
        item.font.name = style["font_latin"]
        item.font.size = Pt(style["body_size"])
        item.font.color.rgb = RGBColor(0, 0, 0)
        for border in item.element.xpath("./w:pPr/w:pBdr"):
            border.getparent().remove(border)
        fonts = item.element.get_or_add_rPr().get_or_add_rFonts()
        for attr in list(fonts.attrib):
            if attr.endswith("Theme") or attr.endswith("theme"):
                del fonts.attrib[attr]
        item.element.get_or_add_rPr().get_or_add_rFonts().set(
            qn("w:eastAsia"),
            style["font_heading_cjk"]
            if name.startswith("Heading") or name == "Title"
            else style["font_body_cjk"],
        )
    document.styles["Title"].font.size = Pt(style["title_size"])
    document.styles["Normal"].paragraph_format.line_spacing = style["line_spacing"]
    document.styles["Normal"].paragraph_format.space_after = Pt(8)
    for level, size in enumerate(style["heading_sizes"], 1):
        document.styles[f"Heading {level}"].font.size = Pt(size)
        document.styles[f"Heading {level}"].paragraph_format.keep_with_next = True
    section = document.sections[0]
    section.page_width, section.page_height = Mm(210), Mm(297)
    for key in ("top", "bottom", "left", "right"):
        setattr(section, f"{key}_margin", Mm(style[f"margin_{key}"]))
    width_mm = 210 - style["margin_left"] - style["margin_right"]
    header = section.header.paragraphs[0]
    header.text = meta["title"]
    header.style = document.styles["Header"]
    header.runs[0].font.size = Pt(9)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    zh = meta.get("language") == "zh"
    footer.add_run("第 " if zh else "Page ")
    field(footer, "PAGE", "1")
    footer.add_run(" 页 共 " if zh else " of ")
    field(footer, "NUMPAGES", "1")
    if zh:
        footer.add_run(" 页")
    for run in footer.runs:
        run.font.size = Pt(9)
    update = OxmlElement("w:updateFields")
    update.set(qn("w:val"), "true")
    document.settings.element.append(update)
    title = document.add_paragraph(meta["title"], style="Title")
    if meta.get("cover", style["cover"]):
        title.paragraph_format.space_before = Mm(45)
    for key in ("subtitle", "author", "date", "classification"):
        if meta.get(key):
            document.add_paragraph(meta[key], style="Subtitle" if key == "subtitle" else "Normal")
    if meta.get("cover", style["cover"]):
        document.add_page_break()
    if meta.get("toc", style["toc"]):
        document.add_paragraph("Contents", style="TOC Heading")
        paragraph = document.add_paragraph()
        # Cached navigation remains readable until Word or LibreOffice updates page numbers.
        begin = OxmlElement("w:fldChar")
        begin.set(qn("w:fldCharType"), "begin")
        paragraph.add_run()._r.append(begin)
        instruction = OxmlElement("w:instrText")
        instruction.text = ' TOC \\o "1-3" \\h \\z \\u '
        paragraph.add_run()._r.append(instruction)
        separate = OxmlElement("w:fldChar")
        separate.set(qn("w:fldCharType"), "separate")
        paragraph.add_run()._r.append(separate)
        for block in model["blocks"]:
            if block["type"] == "heading" and block["level"] <= 3:
                paragraph = document.add_paragraph()
                paragraph.paragraph_format.left_indent = Mm((block["level"] - 1) * 5)
                hyperlink(
                    paragraph, block["anchor"], block["number"] + block["text"], internal=True
                )
        end = OxmlElement("w:fldChar")
        end.set(qn("w:fldCharType"), "end")
        paragraph.add_run()._r.append(end)
        document.add_page_break()
    for index, block in enumerate(model["blocks"], 1):
        kind = block["type"]
        if kind == "heading":
            paragraph = document.add_paragraph(style=f"Heading {block['level']}")
            paragraph.add_run(block["number"])
            inline_docx(paragraph, block["children"], model, width_mm)
            bookmark(paragraph, block["anchor"], index)
        elif kind == "paragraph":
            paragraph = document.add_paragraph()
            if index > 1 and model["blocks"][index - 2]["type"] == "table":
                paragraph.paragraph_format.space_before = Pt(8)
            if block["list_depth"]:
                paragraph.style = (
                    "List Number" if block["list_kind"] == "ordered" else "List Bullet"
                )
                if block["children"] and re.match(r"^\[[ xX]\]", block["children"][0].content):
                    paragraph.style = "Normal"
                paragraph.paragraph_format.left_indent = Mm(5 * block["list_depth"])
            if block["quote"]:
                paragraph.paragraph_format.left_indent = Mm(8)
                paragraph.style = "Quote"
            inline_docx(paragraph, block["children"], model, width_mm)
        elif kind == "table":
            caption = document.add_paragraph(style="Caption")
            caption.paragraph_format.keep_with_next = True
            caption.add_run("表 " if zh else "Table ")
            field(caption, "SEQ Table \\* ARABIC", str(block["number"]))
            bookmark(caption, block["anchor"], index)
            caption.add_run(" — " + block["caption"] if block["caption"] else "")
            table = document.add_table(rows=len(block["rows"]), cols=len(block["rows"][0]))
            table.autofit = True
            border = OxmlElement("w:tblBorders")
            for edge in ("top", "bottom", "left", "right", "insideH", "insideV"):
                element = OxmlElement(f"w:{edge}")
                element.set(qn("w:val"), "single" if edge in {"top", "bottom"} else "nil")
                element.set(qn("w:sz"), "8")
                element.set(qn("w:color"), "666666")
                border.append(element)
            table._tbl.tblPr.append(border)
            lengths = [
                max(len("".join(t.content for t in row[i])) for row in block["rows"])
                for i in range(len(block["rows"][0]))
            ]
            weights = [max(5, min(40, value)) for value in lengths]
            for row_index, row in enumerate(block["rows"]):
                table.rows[row_index]._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
                if row_index == 0:
                    repeat = OxmlElement("w:tblHeader")
                    table.rows[0]._tr.get_or_add_trPr().append(repeat)
                for col, children in enumerate(row):
                    cell = table.cell(row_index, col)
                    cell.width = Mm(width_mm * weights[col] / sum(weights))
                    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                    paragraph = cell.paragraphs[0]
                    paragraph.paragraph_format.space_after = Pt(5)
                    paragraph.paragraph_format.space_before = Pt(5)
                    paragraph.paragraph_format.line_spacing = 1.15
                    inline_docx(paragraph, children, model, width_mm / len(row))
                    if row_index == 0:
                        shading = OxmlElement("w:shd")
                        shading.set(qn("w:fill"), "F1F4F5")
                        cell._tc.get_or_add_tcPr().append(shading)
                        borders = OxmlElement("w:tcBorders")
                        bottom = OxmlElement("w:bottom")
                        bottom.set(qn("w:val"), "single")
                        bottom.set(qn("w:sz"), "4")
                        borders.append(bottom)
                        cell._tc.get_or_add_tcPr().append(borders)
                        for run in paragraph.runs:
                            run.bold = True
        elif kind in {"figure", "math"}:
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.space_before = Pt(8)
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.keep_with_next = kind == "figure"
            add_image(
                paragraph,
                block["data"],
                width_mm if kind == "figure" else width_mm * 0.7,
                alt=block.get("caption", block.get("latex", "")),
            )
            if kind == "figure":
                caption = document.add_paragraph(style="Caption")
                caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
                caption.add_run("图 " if zh else "Figure ")
                field(caption, "SEQ Figure \\* ARABIC", str(block["number"]))
                bookmark(caption, block["anchor"], index)
                caption.add_run(" — " + block["caption"])
        elif kind == "code":
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.line_spacing = 1.1
            paragraph.paragraph_format.keep_together = True
            run = paragraph.add_run(block["text"])
            run.font.name, run.font.size = "Courier New", Pt(9)
    add_footnotes(document, model)
    document.save(path)
